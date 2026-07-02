"""Rebuild the MSFT / GOOGL / PEP workbook with the data grouped by financial
statement (Income Statement, Balance Sheet, Cash Flow, Market) plus the 11
ratios, reusing the shared writer in :mod:`financial_analysis`.

It reads the *original* hand-entered workbook so no figure is lost, evaluates
the simple run-rate / market-cap formulas to plain numbers, normalizes MSFT
from thousands to millions so all three sheets share one unit, and writes a
clean model where every derived line and ratio is a live Excel formula.

    python analyses/scripts/reorganize_existing.py
"""

from __future__ import annotations

import ast
import operator
import re
from pathlib import Path

import openpyxl

from financial_analysis import SheetData, fetch_yfinance, write_ticker_sheet

# line items the original workbook lacks; back-filled from Yahoo Finance so the
# new debt / shareholder-return ratios compute instead of showing 0%.
ENRICH_KEYS = ["Total debt", "Net debt", "Cash dividend paid",
               "Repurchase of capital stock"]

SRC = Path(__file__).resolve().parents[1] / "Analisis_MSFT_GOOGL_PEP.xlsx"
OUT = Path(__file__).resolve().parents[1] / "Analisis_MSFT_GOOGL_PEP_v2.xlsx"

# original row -> line-item key (only INPUT rows we carry over)
ROWS = {
    5: "Revenue", 6: "EBITDA", 7: "EBIT", 8: "Net income", 9: "D&A",
    10: "Total assets", 11: "Cash & mktb sec.", 12: "Non-mktb sec.",
    13: "Current liabilities", 15: "Stockholders' equity", 16: "Total shares",
    17: "Market price", 19: "Total liabilities", 20: "CapEx", 21: "FCF",
    22: "Operating CF", 34: "DY",
}
COLS = ["D", "E", "F", "G"]
PERIOD_LABELS = ["2023", "2024", "2025", "2026-03 (run-rate)"]

# items that are NOT currency (must not be rescaled when converting units)
NON_CURRENCY = {"Market price", "DY"}

_OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
        ast.Div: operator.truediv, ast.USub: operator.neg}


def _eval(node):
    if isinstance(node, ast.BinOp):
        return _OPS[type(node.op)](_eval(node.left), _eval(node.right))
    if isinstance(node, ast.UnaryOp):
        return _OPS[type(node.op)](_eval(node.operand))
    if isinstance(node, ast.Constant):
        return node.value
    raise ValueError(node)


def cell_value(ws, coord):
    v = ws[coord].value
    if isinstance(v, str) and v.startswith("="):
        expr = v[1:]
        if re.fullmatch(r"[0-9.\-+*/() ]+", expr):   # pure arithmetic -> evaluate
            return _eval(ast.parse(expr, mode="eval").body)
        return None                                   # references another cell -> derived, skip
    return v


def read_observations(ws):
    notes = []
    for r in range(38, ws.max_row + 1):
        v = ws.cell(row=r, column=3).value           # column C
        if isinstance(v, str) and v.strip() and v.strip() != "Observaciones":
            notes.append(v.strip())
    return notes


def build():
    wb_src = openpyxl.load_workbook(SRC, data_only=False)
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    for name in wb_src.sheetnames:
        ws = wb_src[name]
        to_millions = (name == "MSFT")   # MSFT is entered in thousands

        inputs: dict[str, list] = {}
        for row, key in ROWS.items():
            series = []
            for c in COLS:
                v = cell_value(ws, f"{c}{row}")
                if v is not None and to_millions and key not in NON_CURRENCY:
                    v = v / 1000.0
                series.append(v)
            inputs[key] = series

        # back-fill debt / shareholder-return lines from Yahoo Finance
        # (index-aligned: 3 full years + 1 trailing, oldest -> newest)
        enrich_note = ""
        try:
            yf_data = fetch_yfinance(name, years=3)
            for key in ENRICH_KEYS:
                src_vals = yf_data.inputs.get(key) or []
                inputs[key] = [src_vals[i] if i < len(src_vals) else None
                               for i in range(len(PERIOD_LABELS))]
            enrich_note = (" Total debt / net debt / dividends / buybacks "
                           "back-filled from Yahoo Finance.")
        except Exception as exc:   # offline or throttled -> leave blank
            print(f"  [warn] could not enrich {name} from yfinance: {exc}")

        data = SheetData(
            ticker=name,
            inputs=inputs,
            periods=[{"label": lbl, "partial": (i == len(PERIOD_LABELS) - 1)}
                     for i, lbl in enumerate(PERIOD_LABELS)],
            unit_label="US$ millions",
            notes=read_observations(ws),
            source_note="Reorganized from Analisis_MSFT_GOOGL_PEP.xlsx "
                        "(MSFT converted from thousands to millions)." + enrich_note,
        )
        out_ws = wb_out.create_sheet(title=name)
        write_ticker_sheet(out_ws, data)

    wb_out.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
